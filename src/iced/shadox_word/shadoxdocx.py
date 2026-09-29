#!/usr/bin/env python
# -*- coding: utf-8 -*-

# THIS DOCUMENT AND ITS CONTENTS ARE PROPERTY OF ARIANEGROUP.
# IT SHALL NOT BE COMMUNICATED TO ANY THIRD PARTY WITHOUT THE OWNER'S
# WRITTEN CONSENT | ARIANEGROUP SAS - ALL RIGHTS RESERVED.

from uuid import uuid4

from docx import Document
import re

class ShadoxDocxHelper():
    def __init__(self, path):
        self.doc = Document(path)

    def make_right_for_template(self, bad_characters):
        """
            Changes the tags to have a template that fits for us
        """
        for paragraph in self.doc.paragraphs:
            tags = re.findall(r'{%(?:(?!%}).)*|{{(?:(?!}}).)*', paragraph.text)
            for tag in tags:
                new_tag = create_id(tag)
                new_tag = replace_apostrophe(new_tag)
                new_tag = replace_dash_id(new_tag, bad_characters)
                new_tag = change_operations(new_tag)
                new_tag = replace_bad_characters(new_tag, bad_characters, 1)
                paragraph.text = paragraph.text.replace(tag, new_tag)
        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        tags = re.findall(r'{%(?:(?!%}).)*|{{(?:(?!}}).)*', paragraph.text)
                        for tag in tags:
                            new_tag = replace_apostrophe(tag)
                            new_tag = replace_bad_characters(new_tag, bad_characters, 1)
                            paragraph.text = paragraph.text.replace(tag, new_tag)

    def save(self, path):
        self.doc.save(path)


def replace_bad_characters(text, bad_characters, dir=0):
    """
        Replace the characters that are forbidden in jinja2 templating
        (with special handling of spaces)
        dir = 1 is the direction when the function encodes the bad characters with the corresponding integer
        dir = 0 is the direction when the function decodes the bad characters from the corresponding integer
    """
    if dir == 1:
        text = re.sub('<[^>]+>', lambda m: m.group(0).replace(" ", bad_characters[" "]), text)

    tag = text.partition("|")[0]
    for key, value in bad_characters.items():
        if key == " ":
            continue
        if dir == 1:
            tag = tag.replace(key, value)
        else:
            tag = tag.replace(value, key)
        text = text.replace(text.partition("|")[0], tag)

    if dir == 0:
        text = re.sub('<[^>]+>', lambda m: m.group(0).replace(bad_characters[" "], " "), text)

    return text

def replace_dash_id(text, bad_characters):
    """
        Replace the - character in the potential id before its replacement for subtraction
    """
    return re.sub('<[^>]+>', lambda m: m.group(0).replace("-", bad_characters["-"]), text)

def replace_apostrophe(text):
    """
        Word doesn't put in the right apostrophes
    """
    text = text.replace("‘", "'")
    text = text.replace("’", "'")
    return text


def create_id(text):
    """
        Create a unique id for each request to be able to request informations on a parameter more than once
    """
    if "<" in text:
        text = text[:text.index("<") + 1] + str(uuid4().int) + text[text.index("<") + 1:]
    return text


def change_operations(text):
    """
        Create filters that will push the operations onto the second template
    """
    tag = text.partition("|")[0]
    add = re.search('\+(\d.)+', tag)
    substract = re.search('-(\d.)+', tag)
    multiply = re.search('\*(\d.)+', tag)
    divide = re.search('/(\d.)+', tag)
    if add:
        tag = tag.replace(add.group(0), "| add(" + add.group(0)[1:] + ")")
    if substract:
        tag = tag.replace(substract.group(0), "| substract(" + substract.group(0)[1:] + ")")
    if multiply:
        tag = tag.replace(multiply.group(0), "| multiply(" + multiply.group(0)[1:] + ")")
    if divide:
        tag = tag.replace(divide.group(0), "| divide(" + divide.group(0)[1:] + ")")
    text = text.replace(text.partition("|")[0], tag)
    return(text)
